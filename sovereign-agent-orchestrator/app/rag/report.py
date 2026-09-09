import json
import re
from collections import Counter
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches
from openpyxl import load_workbook


def _text(value):
    return '' if value is None else str(value).strip()


def _key(value):
    return re.sub(r'[^a-z0-9]+', ' ', _text(value).lower()).strip()


def _count_values(rows, header, splitter=False):
    counts = Counter()
    if not header:
        return {}
    for row in rows:
        value = _text(row.get(header))
        if not value:
            continue
        values = re.split(r'[,;/&]|\band\b', value, flags=re.IGNORECASE) if splitter else [value]
        counts.update(item.strip() for item in values if item.strip())
    return dict(counts.most_common())


def _find_header(headers, *terms):
    for header in headers:
        normalized = _key(header)
        if all(term in normalized for term in terms):
            return header
    return None


def _unique_headers(values):
    counts = Counter()
    headers = []
    for index, value in enumerate(values, 1):
        base = _text(value) or f'Column {index}'
        counts[base] += 1
        headers.append(base if counts[base] == 1 else f'{base} ({counts[base]})')
    return headers


def analyze_workbook(path):
    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for worksheet in workbook.worksheets:
        rows = list(worksheet.values)
        if not rows:
            sheets.append({'name': worksheet.title, 'rows': 0, 'columns': 0})
            continue
        headers = _unique_headers(rows[0])
        records = [dict(zip(headers, row)) for row in rows[1:] if any(value is not None for value in row)]
        sheet = {'name': worksheet.title, 'rows': len(records), 'columns': len(headers), 'headers': headers}
        if _key(worksheet.title) == 'form responses 1':
            branch = _find_header(headers, 'branch')
            year = _find_header(headers, 'year')
            domain = _find_header(headers, 'domain')
            rating = _find_header(headers, 'rate', 'collaboratively')
            follow = _find_header(headers, 'follow')
            community = _find_header(headers, 'joined', 'community')
            sheet['summary'] = {
                'branch': _count_values(records, branch),
                'year': _count_values(records, year),
                'domains': _count_values(records, domain, splitter=True),
                'collaboration_rating': _count_values(records, rating),
                'followed_instagram': _count_values(records, follow),
                'joined_community': _count_values(records, community),
            }
            ratings = [float(row[rating]) for row in records if rating and isinstance(row.get(rating), (int, float))]
            sheet['summary']['average_collaboration_rating'] = round(sum(ratings) / len(ratings), 2) if ratings else None
        elif _key(worksheet.title) == 'marketing':
            time_headers = [header for header in headers if _key(header) == 'time' or _key(header).startswith('time ')]
            time_slots = Counter()
            for time_header in time_headers:
                time_slots.update(_count_values(records, time_header))
            sheet['summary'] = {'time_slots': dict(time_slots.most_common())}
        sheets.append(sheet)
    return {'source': path.name, 'generated_at': datetime.now().isoformat(timespec='seconds'), 'sheets': sheets}


def _add_counts(document, heading, counts):
    document.add_heading(heading, level=2)
    if not counts:
        document.add_paragraph('No values available.')
        return
    for label, count in counts.items():
        document.add_paragraph(f'{label}: {count}', style='List Bullet')


def write_workbook_report(path, output_dir, analysis=None):
    path = Path(path)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    analysis = analysis or analyze_workbook(path)
    json_path = output_dir / f'{path.stem}_report.json'
    docx_path = output_dir / f'{path.stem}_report.docx'
    json_path.write_text(json.dumps(analysis, indent=2, default=str), encoding='utf-8')

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    document.add_heading('TEAM UAS NMIMS Recruitment Report', level=0)
    document.add_paragraph(f'Source workbook: {analysis["source"]}')
    document.add_paragraph(f'Generated locally: {analysis["generated_at"]}')
    document.add_heading('Executive Summary', level=1)
    for sheet in analysis['sheets']:
        document.add_paragraph(f'{sheet["name"]}: {sheet["rows"]} populated response rows and {sheet["columns"]} columns.', style='List Bullet')

    for sheet in analysis['sheets']:
        summary = sheet.get('summary', {})
        if not summary:
            continue
        document.add_heading(sheet['name'], level=1)
        for label, counts in summary.items():
            if isinstance(counts, dict):
                _add_counts(document, label.replace('_', ' ').title(), counts)
            elif counts is not None:
                document.add_paragraph(f'{label.replace("_", " ").title()}: {counts}')

    document.add_heading('Method and Evidence', level=1)
    document.add_paragraph(
        'This report was generated offline from the workbook using the local RAG ingestion and '
        'structured workbook analysis pipeline. No external services or applicant-level records '
        'are required for these aggregate findings. Searchable evidence is indexed under the '
        f'source name {analysis["source"]}.'
    )
    document.save(docx_path)
    return {'json': str(json_path), 'docx': str(docx_path), 'analysis': analysis}


def write_knowledge_transfer_report(paths, output_dir, extractor):
    paths = [Path(path) for path in paths]
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    sources = []
    for path in paths:
        content = extractor(path)
        sources.append({
            'source': path.name,
            'extension': path.suffix.lower(),
            'characters': len(content),
            'content': content,
        })

    generated_at = datetime.now().isoformat(timespec='seconds')
    analysis = {
        'title': 'Knowledge Transfer Report',
        'generated_at': generated_at,
        'sources': [{key: value for key, value in source.items() if key != 'content'} for source in sources],
    }
    json_path = output_dir / 'knowledge_transfer_report.json'
    docx_path = output_dir / 'knowledge_transfer_report.docx'
    json_path.write_text(json.dumps({**analysis, 'documents': sources}, indent=2, default=str), encoding='utf-8')

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    document.add_heading('Knowledge Transfer Report', level=0)
    document.add_paragraph(f'Generated locally: {generated_at}')
    document.add_heading('Executive Summary', level=1)
    document.add_paragraph(
        f'This transfer package consolidates {len(sources)} source documents into one searchable '
        'reference report. The problem statement supplies the domain context, the coding prompt '
        'defines expected agent behavior, and the README provides project-level operating guidance.'
    )
    document.add_heading('Transfer Checklist', level=1)
    for item in (
        'Understand the problem statement, constraints, stakeholders, and expected outcome.',
        'Review the coding prompt for operating rules, implementation standards, and verification expectations.',
        'Use the README for setup, commands, architecture, and day-to-day project operation.',
        'Search the indexed source documents before making assumptions or changing implementation behavior.',
    ):
        document.add_paragraph(item, style='List Bullet')

    document.add_heading('Source Register', level=1)
    for source in sources:
        document.add_paragraph(f'{source["source"]} ({source["characters"]} extracted characters)', style='List Bullet')

    document.add_heading('Source Evidence', level=1)
    for source in sources:
        document.add_heading(source['source'], level=1)
        document.add_paragraph(f'Citation: {source["source"]}')
        content = source['content'].strip()
        document.add_paragraph(content or 'No extractable text was found in this source.')

    document.add_heading('Method', level=1)
    document.add_paragraph(
        'The report was generated offline. Each source was extracted by the local document '
        'pipeline and indexed in the configured RAG store. No external web service was required.'
    )
    document.save(docx_path)
    return {'json': str(json_path), 'docx': str(docx_path), 'analysis': analysis}