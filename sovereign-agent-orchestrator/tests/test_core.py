from pathlib import Path
from app.policy.engine import Policy,Decision
from app.workspace.manager import Workspace
from app.models.router import ModelRouter
from app.rag.service import RagService

def test_policy_unknown_denied(): assert Policy().check('nope',{}).decision==Decision.DENY
def test_workspace_traversal(tmp_path):
 w=Workspace(tmp_path); w.create('j')
 try:w.safe('j','../../secret')
 except ValueError as e: assert str(e)=='PATH_OUTSIDE_JOB_WORKSPACE'
 else: assert False

def test_router_document(): assert ModelRouter('config/models.yaml').route('summarize inspection report')['task_type']=='document_workflow'
def test_router_multimodal(): assert ModelRouter('config/models.yaml').route('inspect scanned drawing image')['task_type']=='multimodal'

def test_rag_extracts_csv_and_searches_without_ollama(tmp_path):
	source = tmp_path / 'findings.csv'
	source.write_text('finding,status\nfire extinguisher,recertify\n')
	rag = RagService(f'sqlite:///{tmp_path / "rag.db"}')
	assert 'fire extinguisher' in rag.extract(source)
	import asyncio
	indexed = asyncio.run(rag.ingest(source))
	assert indexed['chunks'] == 1
	hits = asyncio.run(rag.search('fire extinguisher'))
	assert hits[0]['source'] == 'findings.csv'

def test_rag_extracts_docx(tmp_path):
	from docx import Document
	source = tmp_path / 'report.docx'
	document = Document()
	document.add_paragraph('Pump P-101 requires inspection.')
	document.save(source)
	rag = RagService(f'sqlite:///{tmp_path / "rag.db"}')
	assert 'Pump P-101' in rag.extract(source)

def test_rag_extracts_pptx(tmp_path):
	from pptx import Presentation
	source = tmp_path / 'briefing.pptx'
	presentation = Presentation()
	slide = presentation.slides.add_slide(presentation.slide_layouts[1])
	slide.shapes.title.text = 'Turbine Overhaul'
	slide.placeholders[1].text = 'Bearing replacement due Q3.'
	presentation.save(source)
	rag = RagService(f'sqlite:///{tmp_path / "rag.db"}')
	extracted = rag.extract(source)
	assert 'Turbine Overhaul' in extracted
	assert 'Bearing replacement due Q3.' in extracted

def test_generate_pdf_tool_writes_readable_pdf(tmp_path):
	from app.tools.registry import ToolRegistry
	workspace = Workspace(tmp_path / 'workspace')
	job_id = 'job'
	workspace.create(job_id)
	tools = ToolRegistry(workspace)
	result = tools.execute(job_id, 'generate_pdf', {
		'filename': 'report.pdf',
		'title': 'Turbine Overhaul Report',
		'sections': [{'heading': 'Findings', 'body': 'Bearing replacement due Q3.'}],
		'citations': ['maintenance_log.csv'],
	})
	output = workspace.safe(job_id, result['path'])
	assert output.read_bytes().startswith(b'%PDF')
	import pymupdf
	doc = pymupdf.open(output)
	text = ''.join(page.get_text() for page in doc)
	doc.close()
	assert 'Turbine Overhaul Report' in text
	assert 'Bearing replacement due Q3.' in text

def test_workbook_report_summarizes_recruitment_data(tmp_path):
	from openpyxl import Workbook
	from app.rag.report import analyze_workbook, write_workbook_report
	source = tmp_path / 'responses.xlsx'
	workbook = Workbook()
	sheet = workbook.active
	sheet.title = 'Form Responses 1'
	sheet.append(['Name', 'Branch', 'Year', 'Domain', 'Rate collaboratively'])
	sheet.append(['A', 'B.Tech CE', '2nd', 'Marketing, Technical', 5])
	sheet.append(['B', 'B.Tech CE', '1st', 'Marketing', 3])
	workbook.save(source)
	analysis = analyze_workbook(source)
	assert analysis['sheets'][0]['rows'] == 2
	assert analysis['sheets'][0]['summary']['domains']['Marketing'] == 2
	result = write_workbook_report(source, tmp_path / 'reports', analysis)
	assert Path(result['docx']).is_file()
	assert Path(result['json']).is_file()

def test_knowledge_transfer_report_combines_sources(tmp_path):
	from app.rag.report import write_knowledge_transfer_report
	sources = [tmp_path / 'problem.md', tmp_path / 'README.md']
	sources[0].write_text('# Problem\nBuild a safe knowledge system.')
	sources[1].write_text('# Setup\nRun the local server.')
	result = write_knowledge_transfer_report(sources, tmp_path / 'reports', lambda path: path.read_text())
	assert Path(result['docx']).is_file()
	assert result['analysis']['sources'][0]['source'] == 'problem.md'

def test_requested_tools_are_scoped_and_auditable(tmp_path):
	from app.tools.registry import ToolRegistry
	from app.rag.tiered import TieredRagService
	workspace = Workspace(tmp_path / 'workspace')
	job_id = 'job'
	workspace.create(job_id)
	input_path = workspace.safe(job_id, 'input/source.md', True)
	input_path.write_text('Contact test@example.com or 9876543210.')
	# Production always routes tool calls through the tier-aware wrapper so a tool
	# invocation cannot silently read across the admin/higher/lower boundary.
	urls = {tier: f'sqlite:///{tmp_path / (tier + "_rag.db")}' for tier in ('admin', 'higher', 'lower')}
	rag = TieredRagService(urls, 'http://localhost:11434', 'nomic-embed-text', 'qwen2.5vl:3b')
	tools = ToolRegistry(workspace, rag)
	identity = {'user_id': 'u1', 'tenant_id': 't1', 'role': 'lower'}
	assert tools.execute(job_id, 'ingest_document', {'path': 'input/source.md', 'identity': identity})['chunks'] == 1
	assert tools.execute(job_id, 'list_sources', {'identity': identity})['sources']
	assert tools.execute(job_id, 'ocr_document', {'path': 'input/source.md'})['text']
	redacted = tools.execute(job_id, 'redact_pii', {'path': 'input/source.md'})
	assert redacted['redactions'] == 2
	assert '[REDACTED_EMAIL]' in workspace.safe(job_id, redacted['path']).read_text()
	assert tools.execute(job_id, 'search_db', {'query': 'SELECT name FROM rag_documents', 'identity': identity})['rows']
	import pytest
	with pytest.raises(ValueError, match='SMTP_NOT_CONFIGURED'):
		tools.execute(job_id, 'send_email', {'to': 'team@example.com', 'subject': 'Draft'})
	assert Policy().check('send_email', {}).decision == Decision.REQUIRE_APPROVAL

def test_conversation_history_is_durable_and_tenant_scoped(tmp_path):
 from app.storage.store import Store
 store = Store(f'sqlite:///{tmp_path / "store.db"}')
 identity = {'user_id': 'u1', 'tenant_id': 't1', 'role': 'user'}
 conversation = store.create_conversation('c1', 't1', 'u1', 'Research')
 assert conversation['title'] == 'Research'
 store.add_message('m1', 'c1', 'user', 'What is in the report?')
 store.add_message('m2', 'c1', 'assistant', 'The report contains findings.', [{'chunk_id': 'x'}])
 assert [message['role'] for message in store.messages('c1')] == ['user', 'assistant']
 assert store.messages('c1')[1]['citations'][0]['chunk_id'] == 'x'
 assert store.conversation('c1', {'user_id': 'u2', 'tenant_id': 't1', 'role': 'user'}) is None

def test_delete_document_removes_indexed_chunks(tmp_path):
 source = tmp_path / 'source.md'
 source.write_text('Retention policy requires annual review.')
 rag = RagService(f'sqlite:///{tmp_path / "rag.db"}')
 import asyncio
 asyncio.run(rag.ingest(source, {'file_id': 'file-1', 'tenant_id': 't1'}))
 assert asyncio.run(rag.search('retention policy'))
 assert rag.delete_document('file-1') == 1
 assert asyncio.run(rag.search('retention policy')) == []

def test_file_sharing_and_attachment_scope_are_enforced(tmp_path):
 from app.storage.store import Store
 store = Store(f'sqlite:///{tmp_path / "store.db"}')
 owner = {'user_id': 'alice', 'tenant_id': 't1', 'role': 'user'}
 reader = {'user_id': 'bob', 'tenant_id': 't1', 'role': 'user'}
 store.register_file('f1', 'alice', 't1', 'a.md', str(tmp_path / 'a.md'), {})
 store.register_file('f2', 'alice', 't1', 'b.md', str(tmp_path / 'b.md'), {})
 assert store.accessible_file_ids(reader) == []
 share = store.share_file('f1', owner, 'bob')
 assert share['shared_with_user_id'] == 'bob'
 assert store.accessible_file_ids(reader) == ['f1']
 rag = RagService(f'sqlite:///{tmp_path / "rag.db"}')
 first = tmp_path / 'a.md'; second = tmp_path / 'b.md'
 first.write_text('private alpha policy'); second.write_text('private beta policy')
 rag.ingest_sync(first, {'file_id': 'f1', 'tenant_id': 't1', 'owner_id': 'alice'})
 rag.ingest_sync(second, {'file_id': 'f2', 'tenant_id': 't1', 'owner_id': 'alice'})
 hits = rag.search_sync('private policy', file_ids=['f1'])
 assert {hit['metadata']['file_id'] for hit in hits} == {'f1'}
 assert store.revoke_share(share['id'], owner)
 assert store.accessible_file_ids(reader) == []

def test_jwt_identity_uses_signed_claims(monkeypatch):
 import jwt
 from app.auth import current_identity
 monkeypatch.setenv('AUTH_MODE', 'jwt')
 monkeypatch.setenv('JWT_SECRET', 'test-secret-with-at-least-32-bytes')
 token = jwt.encode({'sub': 'alice', 'tenant_id': 't1', 'role': 'user', 'clearance': 'internal', 'exp': 4102444800}, 'test-secret-with-at-least-32-bytes', algorithm='HS256')
 identity = current_identity(authorization=f'Bearer {token}')
 assert identity['user_id'] == 'alice'
 assert identity['tenant_id'] == 't1'

def test_local_reranker_marks_and_orders_results(tmp_path):
 rag = RagService(f'sqlite:///{tmp_path / "rag.db"}')
 first = tmp_path / 'first.md'; second = tmp_path / 'second.md'
 first.write_text('inspection interval annual review')
 second.write_text('inspection unrelated note')
 rag.ingest_sync(first, {'file_id': 'f1', 'tenant_id': 't1'})
 rag.ingest_sync(second, {'file_id': 'f2', 'tenant_id': 't1'})
 hits = rag.search_sync('inspection annual review', file_ids=['f1', 'f2'])
 assert hits[0]['metadata']['file_id'] == 'f1'
 assert 'local_rerank' in hits[0]['retrieval_method']

def test_request_limiter_blocks_after_limit():
 from app.operations import RequestLimiter
 limiter = RequestLimiter(1)
 assert limiter.allow('user')
 assert not limiter.allow('user')


def test_router_classifies_the_new_task_types():
 router = ModelRouter('config/models.yaml')
 assert router.classify('write a python function with a unit test') == 'coding'
 assert router.classify('calculate the pump flow rate and show the steps') == 'calculation'
 assert router.classify('profile the attached excel workbook') == 'spreadsheet'
 assert router.classify('build a slide deck briefing the board') == 'presentation'
 assert router.classify('transcribe the scanned drawing') == 'multimodal'
 assert router.classify('what time is it') == 'general'


def test_sandbox_runs_code_and_blocks_network(tmp_path):
 from app.tools.sandbox import run_script
 ok = tmp_path / 'ok.py'
 ok.write_text('print(2 + 2)\n')
 result = run_script(ok, tmp_path / 'run', timeout=10)
 assert result['passed'] and result['exit_code'] == 0 and '4' in result['stdout']

 net = tmp_path / 'net.py'
 net.write_text("import socket\nsocket.create_connection(('1.1.1.1', 80), 2)\n")
 blocked = run_script(net, tmp_path / 'run', timeout=10)
 assert not blocked['passed']


def _orchestrator(tmp_path, model):
 from app.storage.store import Store
 from app.tools.registry import ToolRegistry
 from app.verification.verifier import Verifier
 from app.orchestrator.service import Orchestrator
 workspace = Workspace(tmp_path / 'workspace')
 store = Store(f'sqlite:///{tmp_path / "store.db"}')
 tools = ToolRegistry(workspace)
 return Orchestrator(store, workspace, ModelRouter('config/models.yaml'), Policy(), tools, Verifier(str(workspace.root)), model), store


def _job(task):
 return {'job_id': 'j-' + str(abs(hash(task)) % 10000), 'status': 'queued', 'task': task,
         'user_context': {'user_id': 'u', 'role': 'admin', 'tenant_id': 'default'},
         'attachments': [], 'routing': None, 'plan': [], 'tool_calls': [], 'observations': [],
         'verification': None, 'requires_human_approval': False, 'approval': None,
         'artifacts': [], 'final_answer': None, 'error': None}


def test_coding_task_writes_source_and_runs_it_in_the_sandbox(tmp_path):
 import asyncio
 from app.models.adapter import FakeModel
 orchestrator, store = _orchestrator(tmp_path, FakeModel())
 job = _job('write a python function to add two numbers')
 asyncio.run(orchestrator.run(job))
 done = store.get(job['job_id'])
 assert done['status'] == 'done'
 assert done['routing']['task_type'] == 'coding'
 runs = [o for o in done['observations'] if isinstance(o, dict) and o.get('tool') == 'run_python']
 assert runs and runs[0]['passed'] and runs[0]['exit_code'] == 0
 assert done['verification']['checks']['code_executed'] is True
 assert any(a['name'].endswith('.py') for a in done['artifacts'])


def test_document_task_gets_docx_and_pdf_with_a_real_summary(tmp_path):
 import asyncio
 from app.models.adapter import FakeModel
 orchestrator, store = _orchestrator(tmp_path, FakeModel())
 job = _job('summarize the quarterly maintenance report')
 asyncio.run(orchestrator.run(job))
 done = store.get(job['job_id'])
 assert done['status'] == 'done'
 assert done['routing']['task_type'] == 'document_workflow'
 assert any(a['name'].endswith('.docx') for a in done['artifacts'])
 assert any(a['name'].endswith('.pdf') for a in done['artifacts'])
 # final_answer must be a real summary of the model's response, not the old
 # hardcoded "Completed the requested workflow..." placeholder every job used to get.
 assert done['final_answer'] != 'Completed the requested workflow. The verified artifact is available through the artifact API.'
 assert 'document tools' in done['final_answer']


def test_ingest_cites_the_original_filename_not_the_upload_name(tmp_path):
 source = tmp_path / '0a1b2c3d-4e5f-6071-8293-a4b5c6d7e8f9_inspection.md'
 source.write_text('Extinguisher FE-114 overdue for service.')
 rag = RagService(f'sqlite:///{tmp_path / "rag.db"}')
 rag.ingest_sync(source, {'tenant_id': 't1'})
 hits = rag.search_sync('extinguisher overdue')
 assert hits and hits[0]['source'] == 'inspection.md'
 rag.ingest_sync(source, {'tenant_id': 't2', 'source_name': 'Custom Name.md'})
 assert rag.search_sync('extinguisher overdue', metadata={'tenant_id': 't2'})[0]['source'] == 'Custom Name.md'


def test_require_postgres_guard_rejects_sqlite(monkeypatch):
 import importlib
 from app import config as config_module
 monkeypatch.setenv('REQUIRE_POSTGRES', 'true')
 monkeypatch.setenv('DATABASE_URL', 'sqlite:///./x.db')
 importlib.reload(config_module)
 import pytest
 with pytest.raises(RuntimeError, match='POSTGRES_REQUIRED_FOR_PRODUCTION'):
  config_module.validate_production_database_settings(config_module.settings)
 monkeypatch.delenv('REQUIRE_POSTGRES')
 monkeypatch.delenv('DATABASE_URL')
 importlib.reload(config_module)


def test_network_monitor_reports_airgap_status(monkeypatch):
 monkeypatch.setenv('OLLAMA_NO_CLOUD', 'true')
 monkeypatch.delenv('OLLAMA_CLOUD_ENDPOINT', raising=False)
 from app.network import NetworkMonitor
 status = NetworkMonitor().check()
 assert status['ok'] is True and status['ollama_cloud_disabled'] is True
 monkeypatch.setenv('OLLAMA_CLOUD_ENDPOINT', 'https://cloud.example')
 assert NetworkMonitor().check()['ok'] is False


def test_recent_jobs_are_tenant_and_owner_scoped(tmp_path):
 from app.storage.store import Store
 store = Store(f'sqlite:///{tmp_path / "store.db"}')
 store.save({'job_id': 'a', 'status': 'done', 'task': 'one', 'user_context': {'user_id': 'alice', 'tenant_id': 't1'}})
 store.save({'job_id': 'b', 'status': 'done', 'task': 'two', 'user_context': {'user_id': 'bob', 'tenant_id': 't1'}})
 store.save({'job_id': 'c', 'status': 'done', 'task': 'three', 'user_context': {'user_id': 'carol', 'tenant_id': 't2'}})
 alice = {'user_id': 'alice', 'tenant_id': 't1', 'role': 'lower'}
 admin = {'user_id': 'root', 'tenant_id': 't1', 'role': 'admin'}
 assert {j['job_id'] for j in store.recent_jobs(alice)} == {'a'}
 assert {j['job_id'] for j in store.recent_jobs(admin)} == {'a', 'b'}
 assert all('created_at' in j for j in store.recent_jobs(admin))


def test_ollama_adapter_timeout_scales_with_token_budget(monkeypatch):
 from app.models.adapter import OllamaAdapter
 monkeypatch.delenv('LLM_TIMEOUT_SECONDS', raising=False)
 monkeypatch.setenv('LLM_MAX_TOKENS', '100')
 short = OllamaAdapter('http://localhost:11434', 'x')
 assert short.default_timeout == 180  # floor, not shortened for a tiny budget
 monkeypatch.setenv('LLM_MAX_TOKENS', '3072')
 long = OllamaAdapter('http://localhost:11434', 'x')
 assert long.default_timeout == 3072 // 3 + 120  # scales up for a bigger budget
 monkeypatch.setenv('LLM_TIMEOUT_SECONDS', '999')
 override = OllamaAdapter('http://localhost:11434', 'x')
 assert override.default_timeout == 999  # explicit override wins either way
 monkeypatch.delenv('LLM_MAX_TOKENS', raising=False)
 monkeypatch.delenv('LLM_TIMEOUT_SECONDS', raising=False)


def test_ollama_adapter_reports_a_named_error_on_request_failure(monkeypatch):
 import asyncio
 import httpx
 from app.models.adapter import OllamaAdapter

 class _FailingClient:
  def __init__(self, *a, **k): pass
  async def __aenter__(self): return self
  async def __aexit__(self, *a): return False
  async def post(self, *a, **k): raise httpx.ReadTimeout('')  # stringifies to ''

 monkeypatch.setattr('httpx.AsyncClient', _FailingClient)
 adapter = OllamaAdapter('http://localhost:11434', 'qwen-test')
 try:
  asyncio.run(adapter.chat([{'role': 'user', 'content': 'hi'}]))
  assert False, 'expected a RuntimeError'
 except RuntimeError as exc:
  message = str(exc)
  assert 'qwen-test' in message and 'ReadTimeout' in message  # not just an empty string


def test_verification_failure_triggers_bounded_replanning(tmp_path):
 import asyncio

 class FlakyModel:
  model = 'flaky'
  def __init__(self):
   self.calls = 0
  async def chat(self, messages, tools=None, **kwargs):
   self.calls += 1
   if self.calls == 1:
    return {'content': '```python\nraise SystemExit(1)\n```'}
   return {'content': '```python\nprint("fixed")\n```'}

 model = FlakyModel()
 orchestrator, store = _orchestrator(tmp_path, model)
 job = _job('write a python script that prints a value')
 job['options'] = {'max_iterations': 3}
 asyncio.run(orchestrator.run(job))
 done = store.get(job['job_id'])
 assert model.calls == 2
 assert done['status'] == 'done'
 assert done['iteration'] == 2
